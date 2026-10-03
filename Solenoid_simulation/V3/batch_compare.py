#Importing libraries

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
import magpylib as magpy

from Solenoid import create_solenoid
from magnet import create_magnet
from analysis import sweep_axial_force, plot_force_overlay
from electrical import coil_resistance
from recalc_xlsx import recalc as recalc_xlsx


#Fichier d'entree/sortie : une ligne = une configuration de bobine (onglet
#Config_Coil), reference a un aimant reutilisable (onglet Config_Magnet). Voir
#coil_configs.xlsx pour le format attendu (colonnes, exemples). Ce script LIT
#les onglets Config_Coil et Config_Magnet et ECRIT dans l'onglet Electrique
#(colonne force_dispo_maintien_N), car cette valeur vient d'une simulation
#magnetique (magpylib) qu'Excel ne peut pas calculer par une simple formule --
#tout le reste de l'onglet Electrique (regime PWM/maintien inclus) reste des
#formules Excel normales qui se recalculent toutes seules a partir de cette
#valeur injectee.
#
#Chemin relatif au script (et non un chemin absolu propre a une machine) pour
#que le projet fonctionne tel quel une fois clone depuis Github, chez
#n'importe qui.
INPUT_FILE = Path(__file__).resolve().parent / "coil_configs.xlsx"

#Ligne Excel (1-indexee) de la premiere config, ligne d'en-tetes, feuilles.
#Meme convention (titre L1, description, en-tetes L5, donnees a partir de L6)
#sur TOUS les onglets du classeur.
FIRST_DATA_ROW = 6
HEADER_ROW0 = 4    # pour pandas (0-indexe) = ligne Excel 5
ELEC_SHEET = "Electrique"
COIL_SHEET = "Config_Coil"
MAGNET_SHEET = "Config_Magnet"
HOLD_POSITION_CELL = "L2"   # "Position de maintien (mm)" dans l'onglet Electrique
FORCE_HOLD_COL = 9          # colonne I = force_dispo_maintien_N (Electrique)

#Pas d'echantillonnage du balayage. Un peu plus grossier que dans main.py
#(1mm au lieu de 0.5mm) pour que comparer plusieurs configs reste rapide ;
#resserre si tu as besoin de plus de precision sur la config retenue.
SWEEP_STEP = 1e-3   #m

REQUIRED_COIL_COLUMNS = [
    "config_name", "R_mm", "L_mm", "layers", "wire_diameter_mm", "wire_copper_diameter_mm",
    "packing_factor", "turns", "voltage_V", "sens_courant", "magnet_config",
    "empty_bobbin_weight_g", "quantity",
]
REQUIRED_MAGNET_COLUMNS = [
    "magnet_config", "shape", "grade", "Br_T", "dim1_mm", "dim2_mm", "height_mm",
    "x0_mm", "y0_mm", "z0_mm",
]

ATTIRE_VALUES = {"attire", "attiré", "attraction", "attract"}
REPOUSSE_VALUES = {"repousse", "repoussé", "repulsion", "repel"}


def ensure_recalculated(path):
    """Tente de recalculer le classeur via LibreOffice avant de le lire, pour
    eviter d'obliger l'utilisateur a l'ouvrir manuellement dans Excel/LibreOffice
    au prealable. Non bloquant : si LibreOffice est indisponible, on continue
    quand meme (les valeurs en cache, si elles existent, seront utilisees)."""
    print("Recalcul du classeur avant lecture (LibreOffice)...")
    result = recalc_xlsx(str(path), timeout=90)
    if result.get("status") == "success":
        print("  -> recalcul OK.\n")
    else:
        print(f"  -> recalcul non effectue ({result.get('error', 'raison inconnue')}).\n"
              f"     Si les colonnes calculees (turns, wire_copper_diameter_mm, Br_T...) "
              f"apparaissent vides ci-dessous, ouvre {path} dans Excel/LibreOffice, "
              f"enregistre-le une fois, puis relance ce script.\n")


def load_coil_configs(path):
    """Lit le fichier Excel de configurations de bobines (feuille 'Config_Coil',
    en-tetes ligne 5). Garde l'index pandas d'origine (PAS de reset_index) pour
    pouvoir retrouver la ligne Excel exacte de chaque config lors de l'ecriture
    des resultats."""
    df = pd.read_excel(path, sheet_name=COIL_SHEET, header=HEADER_ROW0)

    missing = [c for c in REQUIRED_COIL_COLUMNS if c not in df.columns]
    if missing:
        raise RuntimeError(
            f"Colonnes manquantes dans l'onglet '{COIL_SHEET}' : {missing}.\n"
            f"Colonnes trouvees : {df.columns.tolist()}\n"
            f"-> Le fichier {path} ne correspond probablement pas a la derniere version "
            f"de ce script (coil_configs.xlsx et batch_compare.py doivent venir du meme "
            f"projet). Verifie que tu utilises bien le dernier fichier Excel fourni."
        )

    df = df.dropna(subset=["config_name"])

    #Verification que les colonnes-formules Excel (turns, wire_copper_diameter_mm)
    #contiennent bien des nombres et pas des cellules vides -- ca arrive si le
    #fichier n'a jamais ete recalcule.
    for col in ("turns", "wire_copper_diameter_mm"):
        bad_rows = df[df[col].isna()]
        if len(bad_rows) > 0:
            noms = bad_rows["config_name"].tolist()
            raise RuntimeError(
                f"La colonne '{col}' est vide pour : {noms}.\n"
                f"-> Ouvre {path} dans Excel (ou LibreOffice), verifie que les cellules "
                f"correspondantes de l'onglet {COIL_SHEET} affichent bien un nombre "
                f"(pas vide/erreur), enregistre, puis relance ce script."
            )

    return df


def load_magnet_configs(path):
    """Lit le fichier Excel de configurations d'aimants (feuille 'Config_Magnet',
    en-tetes ligne 5). Indexe par 'magnet_config' pour permettre a plusieurs
    lignes de Config_Coil de partager le meme aimant."""
    #usecols="A:K" : n'importe pas la table de reference Grade/Br (colonnes N:Q,
    #voir Config_Magnet) dans ce dataframe -- elle n'est utile qu'a la formule
    #Excel de Br_T, pas au script.
    df = pd.read_excel(path, sheet_name=MAGNET_SHEET, header=HEADER_ROW0, usecols="A:K")

    missing = [c for c in REQUIRED_MAGNET_COLUMNS if c not in df.columns]
    if missing:
        raise RuntimeError(
            f"Colonnes manquantes dans l'onglet '{MAGNET_SHEET}' : {missing}.\n"
            f"Colonnes trouvees : {df.columns.tolist()}"
        )

    df = df.dropna(subset=["magnet_config"])

    bad_rows = df[df["Br_T"].isna()]
    if len(bad_rows) > 0:
        noms = bad_rows["magnet_config"].tolist()
        raise RuntimeError(
            f"La colonne 'Br_T' est vide pour : {noms}.\n"
            f"-> Verifie que 'grade' est bien une valeur de la table de reference (ou que "
            f"Br_override_T est rempli), ouvre {path} dans Excel/LibreOffice, enregistre-le, "
            f"puis relance ce script."
        )

    if df["magnet_config"].duplicated().any():
        dups = df.loc[df["magnet_config"].duplicated(), "magnet_config"].tolist()
        raise RuntimeError(
            f"magnet_config en double dans '{MAGNET_SHEET}' : {dups}. Chaque nom d'aimant "
            f"doit etre unique (plusieurs bobines peuvent le reference, mais il ne doit "
            f"apparaitre qu'une fois dans Config_Magnet)."
        )

    return df.set_index("magnet_config")


def read_hold_position_mm(path):
    """Lit la constante 'Position de maintien (mm)' (valeur simple, pas une
    formule -> pas besoin de recalcul prealable pour la lire)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    val = wb[ELEC_SHEET][HOLD_POSITION_CELL].value
    return float(val) if val is not None else 10.0


def resolve_current_sign(sens_courant, config_name):
    """'Attire' -> +1 (comportement historique : aimant aimante +Z, attire vers le
    centre de la bobine a courant positif). 'Repousse' -> -1 (meme aimant, courant
    inverse -> force inversee a chaque position)."""
    key = str(sens_courant).strip().lower()
    if key in ATTIRE_VALUES:
        return 1.0
    if key in REPOUSSE_VALUES:
        return -1.0
    raise ValueError(
        f"Config_Coil : sens_courant={sens_courant!r} invalide pour '{config_name}' "
        f"(attendu 'Attire' ou 'Repousse')."
    )


def resolve_magnet_row(magnet_row):
    """Convertit une ligne de Config_Magnet (mm) en parametres magpylib (m)."""
    # bracket notation obligatoire ici : magnet_row.shape est l'attribut pandas
    # Series.shape (une tuple de dimensions), pas la colonne "shape" du tableau.
    shape = str(magnet_row["shape"]).strip().lower()
    dim1 = magnet_row.dim1_mm * 1e-3
    dim2 = (magnet_row.dim2_mm * 1e-3) if pd.notna(magnet_row.dim2_mm) else None
    height = magnet_row.height_mm * 1e-3
    Br = float(magnet_row.Br_T)
    position = (magnet_row.x0_mm * 1e-3, magnet_row.y0_mm * 1e-3, magnet_row.z0_mm * 1e-3)
    return shape, dim1, dim2, height, Br, position


def run_config(coil_row, magnet_defs):
    """Construit le solenoide + l'aimant d'une ligne de config (a pleine tension,
    courant impose par la loi d'Ohm, sens impose par sens_courant), puis balaie la
    force le long de l'axe de la bobine depuis la position de depart de l'aimant."""

    magnet_key = coil_row.magnet_config
    if magnet_key not in magnet_defs.index:
        raise RuntimeError(
            f"Config_Coil!magnet_config={magnet_key!r} (config '{coil_row.config_name}') "
            f"introuvable dans l'onglet {MAGNET_SHEET}. Aimants disponibles : "
            f"{magnet_defs.index.tolist()}."
        )
    magnet_row = magnet_defs.loc[magnet_key]
    shape, dim1, dim2, height, Br, position = resolve_magnet_row(magnet_row)

    R = coil_row.R_mm * 1e-3
    L = coil_row.L_mm * 1e-3
    layers = int(coil_row.layers)
    diameter = coil_row.wire_diameter_mm * 1e-3
    turns = int(coil_row.turns)
    voltage = coil_row.voltage_V

    #Diametre du cuivre NU (sans email) pour la resistance, distinct du diametre
    #total (row.wire_diameter_mm) qui sert a l'encombrement/tours. Desormais
    #calcule automatiquement par Excel (Config_Coil!wire_copper_diameter_mm =
    #wire_diameter_mm - 0.03mm par defaut, surchargeable a la main) -- on relit
    #simplement la valeur telle quelle, peu importe qu'elle vienne de la formule
    #ou d'une saisie manuelle de calibration (onglet Mesures).
    copper_diameter = coil_row.wire_copper_diameter_mm * 1e-3

    #Meme formule que l'onglet Electrique (colonne resistance_ohm) -- les deux
    #DOIVENT rester coherents, voir electrical.py.
    resistance = coil_resistance(R, L, layers, diameter, turns, copper_diameter)
    current_magnitude = voltage / resistance   # loi d'Ohm : le courant n'est pas un choix libre

    sign = resolve_current_sign(coil_row.sens_courant, coil_row.config_name)
    current_signed = sign * current_magnitude

    #Resolution du solenoide adaptee au nombre de tours (~15 points/tour, avec un
    #plancher) : voir la note ajoutee dans Solenoid.py.
    resolution = max(300, int(15 * turns))

    solenoid, _ = create_solenoid(R, L, layers, diameter, turns, current_signed,
                                   plotting=False, resolution=resolution)

    magnet = create_magnet(shape, dim1, dim2, height, Br, position=position,
                            plotting=False)

    #Balayage le long de Z depuis la position de depart de l'aimant (Config_Magnet)
    #jusqu'au centre de la bobine (z_end=0). x0/y0 restent fixes pendant le
    #balayage -- voir analysis.sweep_axial_force.
    z_values, F = sweep_axial_force(solenoid, magnet, start_position=position,
                                     z_end=0.0, step=SWEEP_STEP)

    return {"name": str(coil_row.config_name), "z_values": z_values, "F": F, "z_rest": position[2],
            "resistance": resistance, "current_full": current_magnitude, "sens_courant": coil_row.sens_courant,
            "magnet_config": magnet_key}


def force_at_displacement(res, displacement_mm):
    """Interpole la force (N, magnitude) a un deplacement donne (mm) depuis la
    position de repos, a partir de la courbe deja calculee. La force scale
    exactement lineairement avec le courant (bobine a air, pas de saturation),
    donc PAS besoin de refaire une simulation par courant : ce point (calcule a
    pleine tension) sert de reference pour l'onglet Electrique, qui applique
    lui-meme le duty cycle du PWM."""
    disp = (res["z_rest"] - res["z_values"]) * 1000
    return abs(np.interp(displacement_mm, disp, res["F"][:, 2]))


if __name__ == "__main__":

    if not INPUT_FILE.exists():
        sys.exit(f"Fichier introuvable : {INPUT_FILE}\n"
                  f"(place coil_configs.xlsx dans le meme dossier que batch_compare.py, "
                  f"ou modifie INPUT_FILE en haut du script)")

    ensure_recalculated(INPUT_FILE)

    coil_configs = load_coil_configs(INPUT_FILE)
    magnet_defs = load_magnet_configs(INPUT_FILE)
    print(f"{len(coil_configs)} configuration(s) de bobine chargee(s) depuis {INPUT_FILE}")
    print(f"{len(magnet_defs)} configuration(s) d'aimant chargee(s) (onglet {MAGNET_SHEET})\n")
    print(f"Colonnes lues (bobine) : {coil_configs.columns.tolist()}\n")
    print(f"Premiere config : {coil_configs.iloc[0].to_dict()}\n")

    hold_position_mm = read_hold_position_mm(INPUT_FILE)
    print(f"Position de maintien (depuis {ELEC_SHEET}!{HOLD_POSITION_CELL}) : {hold_position_mm} mm\n")

    results = []
    excel_rows = []   # ligne Excel exacte de chaque config, pour l'ecriture des resultats

    for idx, row in coil_configs.iterrows():
        print(f"-> calcul en cours : {row.config_name} (aimant '{row.magnet_config}', "
              f"sens '{row.sens_courant}') ...")
        res = run_config(row, magnet_defs)
        results.append(res)
        excel_rows.append(idx + FIRST_DATA_ROW)   # idx=0 -> ligne Excel 6, idx=1 -> ligne 7, etc.

    fig = plot_force_overlay(results, zoom_mm=10,
                              title="Comparaison des configurations de bobine")

    #Resume rapide en console
    print("\nResume (force max sur le premier centimetre de course) :")
    for res in results:
        disp_mm = (res["z_rest"] - res["z_values"]) * 1000
        mask = disp_mm <= 10
        idx = np.argmin(np.where(mask, res["F"][:, 2], 0)) if res["sens_courant"].strip().lower() in ATTIRE_VALUES \
            else np.argmax(np.where(mask, res["F"][:, 2], 0))
        print(f"  {res['name']:<28s} R={res['resistance']:.3f}ohm  I_plein={res['current_full']:.2f}A  "
              f"aimant={res['magnet_config']:<14s} Fz_max_1cm = {res['F'][idx,2]:+.4f} N  "
              f"a {disp_mm[idx]:.1f} mm de deplacement")

    #Injection de la force disponible a la position de maintien dans l'onglet
    #Electrique (colonne force_dispo_maintien_N), pour que ses formules
    #PWM/maintien basculent de l'estimation rapide vers la vraie simulation.
    print(f"\nInjection de la force a {hold_position_mm}mm dans '{ELEC_SHEET}' (colonne force_dispo_maintien_N)...")
    wb = openpyxl.load_workbook(INPUT_FILE)
    we = wb[ELEC_SHEET]
    for res, excel_row in zip(results, excel_rows):
        f_hold = force_at_displacement(res, hold_position_mm)
        we.cell(row=excel_row, column=FORCE_HOLD_COL, value=f_hold)
        print(f"  {res['name']:<28s} force dispo a {hold_position_mm}mm = {f_hold:.4f} N  ({f_hold/9.81*1000:.1f} g)")
    wb.save(INPUT_FILE)

    #Recalcul des formules Excel (necessite LibreOffice, voir recalc_xlsx.py). Si
    #indisponible sur cette machine, le fichier se recalculera de toute facon
    #automatiquement a la prochaine ouverture dans Excel/LibreOffice -- ce n'est
    #donc pas bloquant, juste un confort pour voir les resultats a jour immediatement.
    result = recalc_xlsx(str(INPUT_FILE), timeout=90)
    if result.get("status") == "success":
        print(f"\n{INPUT_FILE} recalcule : les colonnes de maintien (PWM) sont a jour.")
    else:
        print(f"\nRecalcul automatique indisponible ({result.get('error', 'raison inconnue')}). "
              f"Ouvre {INPUT_FILE} dans Excel/LibreOffice : les formules se mettront a jour "
              f"automatiquement a l'ouverture.")

    fig.show(renderer="browser")
